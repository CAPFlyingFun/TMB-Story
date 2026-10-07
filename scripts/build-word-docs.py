#!/usr/bin/env python3
"""Rebuild Joshua's Word manuscripts from the chapter files.

    python3 scripts/build-word-docs.py            # every group
    python3 scripts/build-word-docs.py 13-15      # one group
    python3 scripts/build-word-docs.py --zip 21   # also zip groups 01-03 .. 19-21

Three chapters per document, grouped by number (01-03, 04-06, 07-09, ...), named like
Joshua's own files: "13-15 Chapters - TRADDOMIUM Micro Battle.docx". A group that is only
partly written (say just Chapter 13) is still written, with the chapters that exist, and
fills up as the rest arrive. Each file is built on Joshua's own template -- the styles and
page setup of an existing file in manuscripts/ -- with a Title paragraph, "Chapters 13-15",
a status line drawn from the chapters' review_status, then a Heading 1 per chapter and the
prose, one paragraph per manuscript paragraph.

The chapter files are the source; these are exports. Rebuild a group whenever any of its
chapters changes, so manuscripts/ never disagrees with chapters/.
"""

import glob
import os
import re
import shutil
import sys
import zipfile

import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "manuscripts")
TEMPLATE = os.path.join(OUT, "01-03 Chapters - TRADDOMIUM Micro Battle.docx")
NAME = "%02d-%02d Chapters - TRADDOMIUM Micro Battle.docx"


def chapters():
    out = {}
    for path in glob.glob(os.path.join(ROOT, "chapters", "movement-*", "chapter-*.md")):
        text = open(path, encoding="utf-8").read()
        _, fm, body = text.split("---\n", 2)
        n = int(re.search(r"^chapter:\s*(\d+)", fm, re.M).group(1))
        title = re.search(r'^title:\s*"(.*)"', fm, re.M).group(1)
        status = re.search(r"^review_status:\s*(\S+)", fm, re.M).group(1)
        last = re.search(r'^last_approved_revision:\s*"([^"]*)"', fm, re.M)
        paras = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
        paras = [p for p in paras if not p.startswith("# ")]
        out[n] = {"title": title, "status": status, "last": last.group(1) if last else "",
                  "paras": [" ".join(p.split("\n")) for p in paras]}
    return out


def status_line(group):
    states = {c["status"] for c in group.values()}
    if states == {"approved"}:
        return "✅ APPROVED — canon."
    if "revision-pending" in states:
        rev = next((c["last"].split()[0] for c in group.values() if c["last"]), "")
        tail = " Last approved text: revision %s." % rev if rev else ""
        return "🟡 REVISION PENDING SCRIPT REVIEW — not canon until Joshua approves this revision." + tail
    return "🟡 PENDING APPROVAL — not canon until Joshua approves."


def blank_from_template():
    doc = docx.Document(TEMPLATE)
    body = doc.element.body
    for el in list(body):
        if el.tag.endswith("}sectPr"):
            continue
        body.remove(el)
    return doc


def build(a, b, chs):
    group = {n: chs[n] for n in range(a, b + 1) if n in chs}
    if not group:
        return None
    doc = blank_from_template()
    doc.add_paragraph("TRADDOMIUM: Micro Battle!", style="Title")
    doc.add_paragraph("Chapters %d-%d" % (a, b))
    doc.add_paragraph(status_line(group))
    for n in sorted(group):
        doc.add_paragraph("Chapter %d: %s" % (n, group[n]["title"]), style="Heading 1")
        for p in group[n]["paras"]:
            doc.add_paragraph(p)
    path = os.path.join(OUT, NAME % (a, b))
    doc.save(path)
    return path, sorted(group)


def main(argv):
    chs = chapters()
    top = max(chs)
    zip_to = None
    if "--zip" in argv:
        i = argv.index("--zip")
        zip_to = int(argv[i + 1])
        argv = argv[:i] + argv[i + 2:]
    groups = []
    if argv:
        for spec in argv:
            a, b = (int(x) for x in spec.split("-"))
            groups.append((a, b))
    else:
        groups = [(a, a + 2) for a in range(1, top + 1, 3)]
    # The template is itself one of the files being rebuilt, so work from a copy of it.
    tmp = os.path.join(OUT, ".template.docx")
    shutil.copyfile(TEMPLATE, tmp)
    globals()["TEMPLATE"] = tmp
    try:
        for a, b in groups:
            r = build(a, b, chs)
            if r:
                print("wrote %s  (chapters %s)" % (os.path.relpath(r[0], ROOT), ", ".join(map(str, r[1]))))
    finally:
        os.remove(tmp)
    if zip_to:
        zpath = os.path.join(OUT, "TMB Chapters 01-%02d - Word.zip" % zip_to)
        with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
            for a in range(1, zip_to + 1, 3):
                f = os.path.join(OUT, NAME % (a, a + 2))
                z.write(f, os.path.basename(f))
        print("wrote", os.path.relpath(zpath, ROOT))


if __name__ == "__main__":
    main(sys.argv[1:])
